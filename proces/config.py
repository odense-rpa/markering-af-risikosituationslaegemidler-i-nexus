import os
from pathlib import Path
from typing import Optional
from openpyxl import load_workbook
from typing import List
from openpyxl.workbook import Workbook


def load_excel_sheet(filepath: str) -> List[str]:
    
    workbook = load_workbook(filepath)
    worksheet = workbook.active
    
    # Get all data excluding headers (first row)
    data_list = []
    
    # Iterate through all rows starting from row 2 (excluding headers)
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        # Add each cell value to the list, filtering out None values
        for cell_value in row:
            if cell_value is not None:
                # Strip whitespace from the beginning and end of the value
                cleaned_value = str(cell_value).strip()
                if cleaned_value:  # Only add non-empty values after stripping
                    data_list.append(cleaned_value)
    
    return data_list
